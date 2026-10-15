#!/usr/bin/env python3
"""Build an editable one-page FAQ Word document (.docx) with Build to Lead branding.

Usage (prefer build_faq.py):
  python build_faq_docx.py faq.json output.docx [--page a4|letter] [--logo my-logo.png]

The JSON spec is described in SKILL.md. The Word file is editable, so the text can
reflow when people change it; the layout is designed to fit one page as delivered.
If LibreOffice is available, the script checks the page count and shrinks text
until it fits on one page.
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("python-docx is required: pip install python-docx")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from faq_common import BRAND, INSTAGRAM, check_spec, labels, logo_box, resolve_logo, split_placeholders  # noqa: E402

FONT = "Arial"
PT_PER_INCH = 72


def rgb(key: str) -> RGBColor:
    return RGBColor.from_string(BRAND[key])


# --- low-level helpers ----------------------------------------------------

def shade(el, fill_hex: str) -> None:
    """Background color for a cell (tcPr) or a run (rPr)."""
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill_hex)
    pr.append(shd)


def no_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}"); b.set(qn("w:val"), "nil"); borders.append(b)
    tblPr.append(borders)


def cell_border(cell, edge: str, color: str, size_eighths: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    b = OxmlElement(f"w:{edge}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(size_eighths)); b.set(qn("w:color"), color)
    borders.append(b)


def cell_margins(table, top=80, bottom=80, left=120, right=120) -> None:
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)


def para_border(paragraph, edge: str, color: str, size_eighths: int, space: int = 4) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr"); pPr.append(pbdr)
    b = OxmlElement(f"w:{edge}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(size_eighths))
    b.set(qn("w:space"), str(space)); b.set(qn("w:color"), color)
    pbdr.append(b)


def set_widths(table, widths_pt) -> None:
    """Fixed column widths that Word, Google Docs and LibreOffice all respect."""
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths_pt) * 20))); tblW.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_pt):
        gc.set(qn("w:w"), str(int(w * 20)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_pt):
            cell.width = Pt(w)


def cell_pad(cell, left=None, right=None) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, v in (("left", left), ("right", right)):
        if v is not None:
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)


def add_link(paragraph, url: str, text: str, size: float, color: str, bold=False) -> None:
    """Add a clickable hyperlink run."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    col = OxmlElement("w:color"); col.set(qn("w:val"), BRAND.get(color, color)); rpr.append(col)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    if bold:
        rpr.append(OxmlElement("w:b"))
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); run.append(t)
    link.append(run)
    paragraph._p.append(link)


def tight(paragraph, before=0, after=0, line=None) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def add_text(paragraph, text: str, size: float, color: str, bold=False, on_dark=False) -> None:
    """Add text, highlighting UPPERCASE_PLACEHOLDERS."""
    for chunk, is_ph in split_placeholders(str(text)):
        r = paragraph.add_run(chunk)
        r.font.name = FONT; r.font.size = Pt(size); r.bold = bold or is_ph
        if is_ph:
            r.font.color.rgb = rgb("placeholder_fg"); shade(r._r, BRAND["placeholder_bg"])
        else:
            r.font.color.rgb = RGBColor.from_string(color) if len(color) == 6 else rgb(color)


# --- document -------------------------------------------------------------

def build(spec: dict, out: Path, page: str = "a4", logo: str | None = None, scale: float = 1.0) -> None:
    L = labels(spec)
    s = scale
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(9 * s)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    sec = doc.sections[0]
    if page == "letter":
        sec.page_width, sec.page_height = Pt(612), Pt(792)
    else:
        sec.page_width, sec.page_height = Pt(595.3), Pt(841.9)
    M = 34
    sec.left_margin = sec.right_margin = Pt(M)
    sec.top_margin = Pt(M + 4); sec.bottom_margin = Pt(M)
    sec.header_distance = Pt(14); sec.footer_distance = Pt(16)
    inner = sec.page_width.pt - 2 * M

    # header: accent line + logo in the top-right corner (on every page)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tight(hp)
    logo_path = resolve_logo(logo)
    try:
        from PIL import Image
        with Image.open(logo_path) as im:
            lw, lh = logo_box(*im.size)
    except Exception:  # pragma: no cover
        lw, lh = 70, 70
    hp.add_run().add_picture(str(logo_path), width=Pt(lw * 0.92), height=Pt(lh * 0.92))
    para_border(hp, "top", BRAND["teal"], 24, space=6)

    # title block
    p = doc.add_paragraph(); tight(p, after=2); add_text(p, L["eyebrow"], 7.5, "terracotta", bold=True)
    p = doc.add_paragraph(); tight(p, after=1); add_text(p, spec.get("title", ""), 20, "navy", bold=True)
    p = doc.add_paragraph(); tight(p, after=8); add_text(p, spec.get("subtitle", ""), 10, "gray")
    para_border(p, "bottom", BRAND["rule"], 6, space=6)

    # at a glance
    glance = spec.get("at_a_glance", [])[:4]
    if glance:
        t = doc.add_table(rows=1, cols=len(glance)); no_borders(t); cell_margins(t, 70, 70, 120, 80)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, g in zip(t.rows[0].cells, glance):
            shade(cell._tc, BRAND["teal_tint"]); cell_border(cell, "left", BRAND["teal"], 24)
            cp = cell.paragraphs[0]; tight(cp); add_text(cp, str(g.get("label", "")).upper(), 6.5, "terracotta", bold=True)
            cp = cell.add_paragraph(); tight(cp); add_text(cp, g.get("value", ""), min(9 * s, 9.5), "navy", bold=True)
        set_widths(t, [inner / len(glance)] * len(glance))
        sp = doc.add_paragraph(); tight(sp, after=2)

    # FAQs in two columns, balanced by text length, in order
    faqs = spec.get("faqs", [])
    weights = [len(f.get("q", "")) + len(f.get("a", "")) + 60 for f in faqs]
    half, acc, cut = sum(weights) / 2, 0, 0
    for i, w in enumerate(weights):
        if i and acc + w / 2 > half:
            break
        acc += w; cut = i + 1
    cols = [faqs[:cut], faqs[cut:]]
    t = doc.add_table(rows=1, cols=2); no_borders(t); cell_margins(t, 0, 0, 0, 160)
    for cell, items in zip(t.rows[0].cells, cols):
        first = True
        for f in items:
            cp = cell.paragraphs[0] if first else cell.add_paragraph()
            tight(cp, before=0 if first else 6 * s)
            add_text(cp, L["w"].get(f.get("w", ""), "").upper(), 6.5, "teal", bold=True)
            cp = cell.add_paragraph(); tight(cp, after=1); add_text(cp, f.get("q", ""), 9.5 * s, "navy", bold=True)
            cp = cell.add_paragraph(); tight(cp, line=1.1); add_text(cp, f.get("a", ""), 8.8 * s, "gray")
            first = False
    cell_border(t.rows[0].cells[0], "right", BRAND["rule"], 4)
    cell_pad(t.rows[0].cells[1], left=220, right=0)
    set_widths(t, [inner / 2, inner / 2])
    sp = doc.add_paragraph(); tight(sp, after=4)

    # rules of the game
    can, cannot = spec.get("can", []), spec.get("cannot", [])
    if can or cannot:
        t = doc.add_table(rows=1, cols=3); no_borders(t); cell_margins(t, 90, 90, 140, 120)
        for cell, items, title, tint, color, mark in [
            (t.rows[0].cells[0], can, L["can"], "teal_tint", "teal", "✓"),
            (t.rows[0].cells[2], cannot, L["cannot"], "terra_tint", "terracotta", "✗"),
        ]:
            shade(cell._tc, BRAND[tint])
            cp = cell.paragraphs[0]; tight(cp, after=3); add_text(cp, title, 9 * s, color, bold=True)
            for it in items:
                cp = cell.add_paragraph(); tight(cp, after=1.5)
                cp.paragraph_format.left_indent = Pt(12); cp.paragraph_format.first_line_indent = Pt(-12)
                r = cp.add_run(mark + "  "); r.bold = True; r.font.size = Pt(8.6 * s)
                r.font.color.rgb = rgb(color); r.font.name = "Segoe UI Symbol"
                add_text(cp, it, 8.6 * s, "navy")
        set_widths(t, [inner / 2 - 6, 12, inner / 2 - 6])
        sp = doc.add_paragraph(); tight(sp, after=4)

    # help band
    help_ = spec.get("help", {})
    rows = [(L[k], help_[k]) for k in ["contact", "channel", "hours", "response_time", "escalation"] if help_.get(k)]
    if rows:
        t = doc.add_table(rows=1, cols=3); no_borders(t); cell_margins(t, 110, 110, 140, 120)
        c0, c1, c2 = t.rows[0].cells
        for c in (c0, c1, c2):
            shade(c._tc, BRAND["navy"])
        cp = c0.paragraphs[0]; tight(cp); add_text(cp, L["help"], 9, "terracotta", bold=True)
        for cell, chunk in ((c1, rows[::2]), (c2, rows[1::2])):
            first = True
            for k, v in chunk:
                cp = cell.paragraphs[0] if first else cell.add_paragraph(); tight(cp, after=1)
                r = cp.add_run(f"{k}: "); r.bold = True; r.font.size = Pt(8.6 * s); r.font.color.rgb = RGBColor(255, 255, 255)
                add_text(cp, v, 8.6 * s, "FFFFFF")
                first = False
        set_widths(t, [100, (inner - 100) / 2, (inner - 100) / 2])

    # footer: "Made with ♥ by @build2lead.club and @maisfloro.tech" + meta, on every page
    ft = sec.footer.add_table(rows=1, cols=2, width=Pt(inner)); no_borders(ft); cell_margins(ft, 60, 0, 0, 0)
    for c in ft.rows[0].cells:
        cell_border(c, "top", BRAND["rule"], 6)
    fp = ft.rows[0].cells[0].paragraphs[0]; tight(fp)
    add_text(fp, L["made"] + " ", 8, "navy", bold=True)
    r = fp.add_run("♥"); r.font.size = Pt(9); r.font.color.rgb = rgb("terracotta"); r.font.name = "Segoe UI Symbol"
    (h1, u1), (h2, u2) = INSTAGRAM
    add_text(fp, " " + L["by"] + " ", 8, "navy", bold=True)
    add_link(fp, u1, h1, 8, "terracotta", bold=True)
    add_text(fp, " " + L["and"] + " ", 8, "navy", bold=True)
    add_link(fp, u2, h2, 8, "terracotta", bold=True)
    mp = ft.rows[0].cells[1].paragraphs[0]; tight(mp); mp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta = "  ·  ".join(f"{L[k]}: {spec[k]}" for k in ["owner", "version", "updated"] if spec.get(k))
    add_text(mp, meta, 7.5, "gray")
    set_widths(ft, [inner * 0.47, inner * 0.53])
    # drop the empty default footer paragraph so the table sits alone
    first = sec.footer.paragraphs[0]
    if not first.text:
        first._p.getparent().remove(first._p)

    doc.core_properties.title = spec.get("title", "FAQ")
    doc.core_properties.author = "Build to Lead"
    doc.save(out)


def page_count(docx_path: Path) -> int | None:
    """Count pages by converting to PDF with LibreOffice, if available."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx_path)],
                           check=True, capture_output=True, timeout=120)
            pdf = Path(tmp) / (docx_path.stem + ".pdf")
            data = pdf.read_bytes()
            return data.count(b"/Type /Page") - data.count(b"/Type /Pages") or None
        except Exception:
            return None


def build_fitted(spec: dict, out: Path, page: str = "a4", logo: str | None = None) -> None:
    for w in check_spec(spec):
        print(f"WARNING: {w}", file=sys.stderr)
    can_measure = bool(shutil.which("soffice") or shutil.which("libreoffice"))
    if not can_measure:
        print("NOTE: LibreOffice not found, so the one-page fit wasn't checked; "
              "fix any warning above before sharing.", file=sys.stderr)
    scales = (1.3, 1.25, 1.2, 1.15, 1.1, 1.05, 1.0, 0.95, 0.9, 0.85, 0.8) if can_measure else (1.0,)
    for scale in scales:
        build(spec, out, page, logo, scale)
        n = page_count(out) if can_measure else None
        if n is None or n <= 1:
            return
    print("WARNING: the Word file is longer than one page; shorten the answers.", file=sys.stderr)


def main() -> int:
    ap = argparse.ArgumentParser(description="Build a one-page FAQ Word document with Build to Lead branding")
    ap.add_argument("spec"); ap.add_argument("output")
    ap.add_argument("--page", choices=["a4", "letter"], default="a4")
    ap.add_argument("--logo", help="PNG/JPG to use instead of the Build to Lead logo")
    a = ap.parse_args()
    spec = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    build_fitted(spec, Path(a.output), a.page, a.logo)
    print(f"Word document created: {a.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
